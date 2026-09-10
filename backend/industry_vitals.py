"""Per-template vitals: unique positive figures + two chart series each, editable per client,
CSV / Excel import. Every one of the 33 templates has its own bespoke entry — nothing is generic."""
import csv
import io
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from fastapi import Depends, File, HTTPException, UploadFile
from pydantic import BaseModel

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
SEASONS = ["Spring", "Summer", "Autumn", "Winter"]


def _labels(kind: str, n: int) -> List[str]:
    if kind == "years":
        return [str(2021 + i) for i in range(n)]
    if kind == "quarters":
        return [f"Q{(i % 4) + 1} '{24 + i // 4}" for i in range(n)]
    if kind == "months":
        return MONTHS[:n]
    if kind == "weeks":
        return [f"Wk {i + 1}" for i in range(n)]
    if kind == "days":
        return DAYS[:n]
    if kind == "seasons":
        return SEASONS[:n]
    return [f"{i + 1}" for i in range(n)]


# key -> (section title, [(label, prefix, value, suffix) x3],
#         (chart1 label, x kind, values), (chart2 label, x kind, values), (variant1, variant2))
T: Dict[str, Any] = {
    "hvac": ("Service results",
             [("First-visit fixes", "", 96, "%"), ("Jobs completed", "", 4820, ""), ("Avg. arrival time", "", 38, " min")],
             ("Jobs completed by month", "months", [280, 305, 340, 372, 355, 410, 468, 490, 430, 398, 372, 415]),
             ("Heating vs cooling mix", "seasons", [46, 88, 52, 74]),
             ("area", "bars")),
    "healthcare": ("Patient outcomes",
                   [("Patient satisfaction", "", 98, "%"), ("Patients seen", "", 12400, ""), ("Avg. wait time", "", 9, " min")],
                   ("Visits per month", "months", [820, 865, 910, 940, 980, 1010, 990, 1035, 1080, 1120, 1160, 1210]),
                   ("Care mix", "", [34, 26, 19, 13, 8]),
                   ("wave", "donut")),
    "construction": ("Project delivery",
                     [("On-time delivery", "", 97, "%"), ("Projects completed", "", 386, ""), ("Safety record", "", 0, " incidents")],
                     ("Projects per quarter", "quarters", [14, 17, 19, 22, 24, 27, 29, 33]),
                     ("Build type split", "", [42, 31, 17, 10]),
                     ("step", "stacked")),
    "fitness": ("Member results",
                [("Goals achieved", "", 93, "%"), ("Active members", "", 2180, ""), ("Avg. weekly visits", "", 4.2, "")],
                ("Member growth", "months", [1180, 1265, 1340, 1410, 1520, 1640, 1710, 1795, 1880, 1965, 2060, 2180]),
                ("Class attendance by day", "days", [78, 84, 91, 88, 95, 72, 58]),
                ("area", "bars")),
    "retail": ("Store performance",
               [("Customer rating", "", 4.9, "/5"), ("Orders shipped", "", 38600, ""), ("Repeat buyers", "", 64, "%")],
               ("Sales by month", "months", [128, 141, 158, 172, 166, 189, 204, 218, 236, 268, 342, 396]),
               ("Category mix", "", [38, 24, 18, 12, 8]),
               ("bars", "donut")),
    "hospitality": ("Guest experience",
                    [("Guest rating", "", 4.8, "/5"), ("Nights booked", "", 18450, ""), ("Repeat guests", "", 57, "%")],
                    ("Occupancy by month", "months", [68, 71, 76, 82, 88, 94, 97, 96, 89, 83, 77, 86]),
                    ("Stay length mix", "", [31, 28, 22, 19]),
                    ("area", "radar")),
    "finance": ("Portfolio results",
                [("Avg. annual return", "", 11.4, "%"), ("Assets managed", "$", 640, "M"), ("Client retention", "", 98, "%")],
                ("Growth by year", "years", [312, 368, 431, 512, 640]),
                ("Allocation", "", [44, 26, 18, 12]),
                ("candles", "donut")),
    "it_services": ("Service levels",
                    [("Uptime", "", 99.98, "%"), ("Tickets resolved", "", 27400, ""), ("Avg. first response", "", 4, " min")],
                    ("Tickets by month", "months", [1840, 1920, 2010, 2140, 2080, 2260, 2340, 2410, 2380, 2450, 2520, 2610]),
                    ("Response SLA", "", [92, 96, 99, 99.98]),
                    ("step", "gauge")),
    "creative_studio": ("Studio impact",
                        [("Client retention", "", 92, "%"), ("Projects shipped", "", 214, ""), ("Awards won", "", 18, "")],
                        ("Launches per quarter", "quarters", [11, 14, 16, 19, 21, 24, 26, 29]),
                        ("Discipline mix", "", [34, 27, 21, 18]),
                        ("bubble", "radar")),
    "logistics": ("Fleet performance",
                  [("On-time delivery", "", 99.2, "%"), ("Shipments moved", "", 184000, ""), ("Fleet utilisation", "", 94, "%")],
                  ("Loads per month", "months", [12800, 13400, 14100, 14800, 15200, 15900, 16400, 16800, 17200, 17600, 18100, 18700]),
                  ("Lane mix", "", [41, 29, 18, 12]),
                  ("area", "stacked")),
    "saas": ("Product metrics",
             [("Uptime", "", 99.99, "%"), ("Active workspaces", "", 8640, ""), ("Avg. response time", "", 120, " ms")],
             ("Weekly active use", "weeks", [4200, 4480, 4760, 5100, 5480, 5820, 6240, 6680, 7120, 7580, 8100, 8640]),
             ("Plan mix", "", [46, 32, 22]),
             ("wave", "stacked")),
    "legal": ("Case results",
              [("Cases won", "", 94, "%"), ("Clients represented", "", 1480, ""), ("Avg. resolution", "", 5.2, " months")],
              ("Matters closed", "quarters", [58, 64, 71, 78, 84, 91, 98, 106]),
              ("Practice area mix", "", [33, 26, 22, 19]),
              ("bars", "radar")),
    "education": ("Student outcomes",
                  [("Graduation rate", "", 99, "%"), ("Students enrolled", "", 1260, ""), ("Avg. class size", "", 14, "")],
                  ("Enrolment growth", "years", [780, 862, 954, 1080, 1260]),
                  ("Programme mix", "", [36, 28, 21, 15]),
                  ("step", "donut")),
    "real_estate": ("Sales performance",
                    [("Homes sold", "", 640, ""), ("Avg. days on market", "", 18, " days"), ("List-to-sale price", "", 102, "%")],
                    ("Closings by quarter", "quarters", [58, 66, 74, 81, 92, 104, 118, 134]),
                    ("Property type mix", "", [42, 27, 19, 12]),
                    ("area", "bubble")),
    "restaurant": ("Kitchen numbers",
                   [("Guest rating", "", 4.9, "/5"), ("Covers served", "", 84200, ""), ("Avg. wait", "", 11, " min")],
                   ("Covers by day", "days", [420, 448, 512, 596, 784, 862, 690]),
                   ("Menu mix", "", [34, 26, 22, 18]),
                   ("bars", "donut")),
    "events": ("Event results",
               [("Attendee rating", "", 4.9, "/5"), ("Guests hosted", "", 64800, ""), ("Events delivered", "", 412, "")],
               ("Events by season", "seasons", [86, 132, 118, 76]),
               ("Format mix", "", [38, 27, 20, 15]),
               ("stacked", "radar")),
    "veterinary": ("Animal care results",
                   [("Recovery rate", "", 97, "%"), ("Patients treated", "", 9840, ""), ("Same-day appointments", "", 88, "%")],
                   ("Visits per month", "months", [640, 682, 714, 760, 812, 868, 904, 942, 908, 876, 848, 892]),
                   ("Species mix", "", [48, 34, 11, 7]),
                   ("wave", "bubble")),
    "dental": ("Smile outcomes",
               [("Treatment success", "", 99, "%"), ("Smiles restored", "", 7420, ""), ("Avg. appointment", "", 42, " min")],
               ("Appointments by month", "months", [520, 548, 586, 612, 648, 674, 702, 736, 758, 784, 806, 842]),
               ("Treatment mix", "", [40, 24, 20, 16]),
               ("area", "gauge")),
    "accounting": ("Books & filings",
                   [("Filed on time", "", 100, "%"), ("Returns prepared", "", 3260, ""), ("Avg. saving found", "$", 4800, "")],
                   ("Filings by quarter", "quarters", [420, 468, 512, 604, 648, 712, 786, 864]),
                   ("Service mix", "", [37, 29, 20, 14]),
                   ("step", "stacked")),
    "landscaping": ("Grounds results",
                    [("Client retention", "", 95, "%"), ("Gardens maintained", "", 1840, ""), ("Avg. visit", "", 96, " min")],
                    ("Visits by season", "seasons", [1240, 2180, 1680, 620]),
                    ("Service mix", "", [42, 26, 19, 13]),
                    ("bars", "bubble")),
    "photography": ("Shoot results",
                    [("Client rating", "", 5.0, "/5"), ("Shoots delivered", "", 1280, ""), ("Avg. turnaround", "", 4, " days")],
                    ("Shoots per quarter", "quarters", [64, 72, 86, 98, 112, 126, 141, 158]),
                    ("Session mix", "", [34, 28, 22, 16]),
                    ("bubble", "donut")),
    "automotive": ("Workshop numbers",
                   [("First-time fix rate", "", 97, "%"), ("Vehicles serviced", "", 14600, ""), ("Avg. turnaround", "", 6, " hrs")],
                   ("Jobs by month", "months", [980, 1024, 1096, 1140, 1188, 1236, 1280, 1324, 1372, 1418, 1466, 1512]),
                   ("Work mix", "", [38, 27, 21, 14]),
                   ("candles", "gauge")),
    "beauty": ("Chair results",
               [("Client rating", "", 4.9, "/5"), ("Appointments", "", 21400, ""), ("Rebooking rate", "", 78, "%")],
               ("Bookings by day", "days", [86, 94, 112, 138, 176, 208, 96]),
               ("Treatment mix", "", [36, 26, 22, 16]),
               ("wave", "radar")),
    "insurance": ("Cover & claims",
                  [("Claims approved", "", 96, "%"), ("Policies in force", "", 18600, ""), ("Avg. claim time", "", 3, " days")],
                  ("Policies by year", "years", [9800, 11600, 13800, 16100, 18600]),
                  ("Cover mix", "", [39, 28, 19, 14]),
                  ("area", "stacked")),
    "pet_grooming": ("Grooming results",
                     [("Happy pets", "", 99, "%"), ("Grooms completed", "", 8640, ""), ("Avg. appointment", "", 68, " min")],
                     ("Grooms by month", "months", [580, 612, 664, 702, 748, 786, 824, 862, 828, 796, 772, 840]),
                     ("Breed size mix", "", [44, 32, 24]),
                     ("bars", "bubble")),
    "hvac_plumbing": ("Callout performance",
                      [("Same-day callouts", "", 94, "%"), ("Callouts resolved", "", 11800, ""), ("Avg. response", "", 46, " min")],
                      ("Callouts by month", "months", [820, 864, 912, 948, 984, 1042, 1108, 1146, 1064, 1008, 962, 1022]),
                      ("Trade split", "", [52, 48]),
                      ("step", "donut")),
    "coworking": ("Community numbers",
                  [("Desk occupancy", "", 92, "%"), ("Members hosted", "", 3480, ""), ("Avg. stay", "", 14, " months")],
                  ("Occupancy by month", "months", [64, 68, 72, 76, 79, 83, 86, 88, 90, 91, 92, 94]),
                  ("Space mix", "", [41, 29, 18, 12]),
                  ("area", "radar")),
    "wellness": ("Wellbeing results",
                 [("Client wellbeing lift", "", 86, "%"), ("Sessions delivered", "", 16400, ""), ("Return rate", "", 81, "%")],
                 ("Sessions by month", "months", [1080, 1146, 1212, 1284, 1352, 1420, 1488, 1556, 1612, 1684, 1748, 1820]),
                 ("Therapy mix", "", [33, 27, 22, 18]),
                 ("wave", "gauge")),
    "cleaning": ("Service standards",
                 [("Quality score", "", 98, "%"), ("Cleans delivered", "", 42800, ""), ("Avg. arrival", "", 12, " min early")],
                 ("Cleans by week", "weeks", [640, 672, 704, 738, 766, 794, 826, 858, 884, 912, 946, 978]),
                 ("Contract mix", "", [46, 31, 23]),
                 ("bars", "stacked")),
    "music_school": ("Student progress",
                     [("Exams passed", "", 98, "%"), ("Students taught", "", 1840, ""), ("Avg. lesson", "", 45, " min")],
                     ("Enrolment by year", "years", [820, 1010, 1240, 1520, 1840]),
                     ("Instrument mix", "", [34, 26, 22, 18]),
                     ("step", "radar")),
    "nonprofit": ("Impact numbers",
                  [("Funds to programmes", "", 92, "%"), ("People supported", "", 68400, ""), ("Volunteer hours", "", 124000, "")],
                  ("People supported by year", "years", [24800, 34200, 44600, 56200, 68400]),
                  ("Programme spend", "", [43, 28, 17, 12]),
                  ("area", "donut")),
    "architecture": ("Practice results",
                     [("Planning approvals", "", 96, "%"), ("Projects designed", "", 268, ""), ("Awards won", "", 12, "")],
                     ("Projects per quarter", "quarters", [12, 15, 18, 21, 24, 26, 29, 32]),
                     ("Sector mix", "", [37, 26, 22, 15]),
                     ("bubble", "stacked")),
    "test_template": ("Sandbox figures",
                      [("Signal quality", "", 99, "%"), ("Runs executed", "", 5400, ""), ("Avg. latency", "", 82, " ms")],
                      ("Runs by week", "weeks", [280, 312, 348, 384, 420, 458, 492, 528, 566, 602, 640, 682]),
                      ("Suite mix", "", [40, 33, 27]),
                      ("candles", "gauge")),
}

FALLBACK = ("By the numbers",
            [("Client satisfaction", "", 98, "%"), ("Projects delivered", "", 640, ""), ("Avg. response time", "", 12, " min")],
            ("Growth by quarter", "quarters", [42, 48, 56, 64, 73, 84, 96, 110]),
            ("Service mix", "", [38, 27, 20, 15]),
            ("area", "donut"))


def spec_for(key: str, industry: str = "") -> Dict[str, Any]:
    title, mets, c1, c2, variants = T.get(key, FALLBACK)
    return {
        "title": title,
        "metrics": [{"label": lab, "prefix": pre, "value": val, "suffix": suf} for lab, pre, val, suf in mets],
        "series_label": c1[0],
        "labels": _labels(c1[1], len(c1[2])),
        "series": list(c1[2]),
        "series2_label": c2[0],
        "labels2": _labels(c2[1], len(c2[2])) if c2[1] else [f"Part {i + 1}" for i in range(len(c2[2]))],
        "series2": list(c2[2]),
        "source_label": "",
        "source_label2": "",
        "variants": list(variants),
        "source": "demo",
    }


def register(api, db, get_current_user, get_user_app, log_activity):

    class VitalsIn(BaseModel):
        title: Optional[str] = None
        series_label: Optional[str] = None
        labels: Optional[List[str]] = None
        series: Optional[List[float]] = None
        series2_label: Optional[str] = None
        labels2: Optional[List[str]] = None
        series2: Optional[List[float]] = None
        source_label: Optional[str] = None
        source_label2: Optional[str] = None
        metrics: Optional[List[Dict[str, Any]]] = None

    def _key(app: Dict[str, Any]) -> str:
        return (app.get("motion_profile") or {}).get("template_key") or app.get("site_niche") or ""

    def _resolved(app: Dict[str, Any]) -> Dict[str, Any]:
        base = spec_for(_key(app), app.get("industry") or "")
        return {**base, **(app.get("vitals") or {})}

    @api.get("/public/vitals/{key}")
    async def public_spec(key: str):
        return spec_for(key)

    @api.get("/apps/{app_id}/vitals")
    async def get_vitals(app_id: str, user: dict = Depends(get_current_user)):
        app = await get_user_app(app_id, user)
        return {**_resolved(app), "demo": spec_for(_key(app), app.get("industry") or "")}

    @api.put("/apps/{app_id}/vitals")
    async def put_vitals(app_id: str, body: VitalsIn, user: dict = Depends(get_current_user)):
        app = await get_user_app(app_id, user)
        patch = {k: v for k, v in body.dict().items() if v is not None}
        cur = {**(app.get("vitals") or {}), **patch, "updated_at": datetime.now(timezone.utc).isoformat()}
        await db.apps.update_one({"app_id": app_id}, {"$set": {"vitals": cur}})
        await log_activity(app_id, user["user_id"], "vitals.save", "Updated the figures shown on the site")
        return _resolved({**app, "vitals": cur})

    @api.post("/apps/{app_id}/vitals/reset")
    async def reset_vitals(app_id: str, user: dict = Depends(get_current_user)):
        app = await get_user_app(app_id, user)
        await db.apps.update_one({"app_id": app_id}, {"$unset": {"vitals": ""}})
        await log_activity(app_id, user["user_id"], "vitals.reset", "Restored the demo figures")
        return _resolved({**app, "vitals": None})

    @api.post("/apps/{app_id}/vitals/import")
    async def import_vitals(app_id: str, file: UploadFile = File(...),
                            target: str = "series", user: dict = Depends(get_current_user)):
        """Two-column sheet (label, value), optional header row. .csv, .xlsx or .xlsm."""
        app = await get_user_app(app_id, user)
        raw = await file.read()
        rows: List[List[Any]] = []
        name = (file.filename or "").lower()
        if name.endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(raw), data_only=True)
                rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            except ImportError:
                raise HTTPException(400, "Please save the sheet as CSV and upload it again")
            except Exception:
                raise HTTPException(400, "That file could not be read as a spreadsheet")
        elif name.endswith(".xls"):
            raise HTTPException(400, "Old .xls files are not supported — save as .xlsx or .csv")
        else:
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8", "ignore"))))
        labels, cols = [], []
        for r in rows:
            if not r or len(r) < 2:
                continue
            nums = []
            for cell in r[1:4]:
                try:
                    nums.append(round(float(str(cell).replace(",", "").replace("%", "").replace("$", "").strip()), 2))
                except (TypeError, ValueError):
                    nums.append(None)
            if nums and nums[0] is None:
                continue    # header rows and blanks are skipped
            lab = str(r[0] if r[0] is not None else "").strip()[:18]
            if not lab:
                continue
            labels.append(lab)
            cols.append(nums)
        if len(cols) < 2:
            raise HTTPException(400, "Need at least two rows of label,value data")
        labels, cols = labels[:12], cols[:12]
        first = [c[0] for c in cols]
        second = [c[1] for c in cols] if all(len(c) > 1 and c[1] is not None for c in cols) else None
        cur = {**(app.get("vitals") or {}), "source": "imported",
               "imported_at": datetime.now(timezone.utc).isoformat(),
               "imported_file": file.filename}
        if target == "series2":
            cur.update({"labels2": labels, "series2": first})
            charts = 1
        else:
            cur.update({"labels": labels, "series": first})
            charts = 1
            if second:      # a third column auto-fills the second chart in the same upload
                cur.update({"labels2": labels, "series2": second})
                charts = 2
        await db.apps.update_one({"app_id": app_id}, {"$set": {"vitals": cur}})
        await log_activity(app_id, user["user_id"], "vitals.import",
                           f"Imported {len(first)} figures from {file.filename}")
        return {"ok": True, "points": len(first), "charts": charts, **_resolved({**app, "vitals": cur})}
