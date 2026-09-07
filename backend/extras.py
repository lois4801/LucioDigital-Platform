import os
import io
import re
import csv
import json
import time
import base64
import asyncio
import logging
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

import requests
import stripe
import dns.resolver
from fastapi import HTTPException, Depends, Request, UploadFile, File
from pydantic import BaseModel, Field

logger = logging.getLogger("agency.extras")

EMERGENT_LLM_KEY = os.environ.get("EMERGENT_LLM_KEY", "")
# Fal Universal Key supports queue inference through Emergent proxy, not Platform APIs.
INTEGRATION_PROXY_BASE = os.environ.get("INTEGRATION_PROXY_URL", "https://integrations.emergentagent.com").rstrip("/")
FAL_CONTROL = f"{INTEGRATION_PROXY_BASE}/api/v1/fal"
FAL_QUEUE_BASE = f"{FAL_CONTROL}/queue"
FAL_VIDEO_ENDPOINT = "fal-ai/minimax/hailuo-02/standard/text-to-video"

stripe.api_key = os.environ.get("STRIPE_SECRET_KEY") or "sk_test_emergent"
STRIPE_WEBHOOK_SECRET = os.environ.get("STRIPE_WEBHOOK_SECRET", "")
TAX_MODE = "full"

DNS_CNAME_TARGET = "tenants.luciostudio.app"


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def uid(prefix):
    import uuid
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


# ---------- Models ----------
class ImageIn(BaseModel):
    prompt: str
    style: str = "marketing"


class VoiceIn(BaseModel):
    text: str
    voice_id: str = "21m00Tcm4TlvDq8ikWAM"


class VideoIn(BaseModel):
    prompt: str
    duration: str = "6"


class ElevenKeyIn(BaseModel):
    api_key: str


class CheckoutIn(BaseModel):
    lookup_key: str
    app_id: str
    origin_url: str


class DomainIn(BaseModel):
    domain: str


class PlanIn(BaseModel):
    name: str
    price: float
    interval: str = "month"
    features: List[str] = []
    description: str = ""


class PlansReplaceIn(BaseModel):
    plans: List[PlanIn]


DEFAULT_PLANS = [
    {"name": "Starter", "price": 29.0, "interval": "month", "description": "Single hosted tenant", "features": ["1 hosted tenant", "Custom domain", "Email support"]},
    {"name": "Pro", "price": 99.0, "interval": "month", "description": "Growing client portfolio", "features": ["10 hosted tenants", "AI Media Studio", "Priority support"]},
    {"name": "Scale", "price": 299.0, "interval": "month", "description": "Agency at scale", "features": ["Unlimited tenants", "Dedicated SLA", "White-label handoff"]},
]


def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


# ---------- Fal helpers (blocking, run in thread) ----------
def _fal_headers(extra=None):
    h = {"Authorization": f"Bearer {EMERGENT_LLM_KEY}", "Content-Type": "application/json"}
    if extra:
        h.update(extra)
    return h


def _trusted_queue_url(value: str) -> str:
    c, b = urlparse(value), urlparse(FAL_QUEUE_BASE)
    if c.scheme != b.scheme or c.netloc != b.netloc or not c.path.startswith(b.path.rstrip("/") + "/"):
        raise RuntimeError("Proxy returned an untrusted media queue URL")
    return value


def run_fal_video(payload: dict, timeout_seconds: int = 600) -> dict:
    r = requests.post(f"{FAL_CONTROL}/proxy", headers=_fal_headers({"X-Fal-Target-Url": f"https://queue.fal.run/{FAL_VIDEO_ENDPOINT}"}),
                      json=payload, timeout=60)
    if r.status_code == 402:
        raise RuntimeError("Insufficient Universal Key credits. Add balance in Profile → Universal Key.")
    r.raise_for_status()
    sub = r.json()
    status_url, response_url = _trusted_queue_url(sub["status_url"]), _trusted_queue_url(sub["response_url"])
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        try:
            s = requests.get(status_url, headers=_fal_headers(), timeout=30)
        except (requests.ConnectTimeout, requests.ReadTimeout, requests.ConnectionError):
            time.sleep(2)
            continue
        s.raise_for_status()
        body = s.json()
        st = (body.get("status") or "").upper()
        if st in {"COMPLETED", "OK"}:
            res = requests.get(response_url, headers=_fal_headers(), timeout=60)
            res.raise_for_status()
            return res.json()
        if st in {"FAILED", "CANCELLED", "CANCELED", "ERROR"}:
            raise RuntimeError(f"Video generation failed: {st}")
        time.sleep(3)
    raise RuntimeError("Timed out waiting for video generation")


# ---------- Pricing document parsing ----------
def _to_price(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"[\d][\d,]*\.?\d*", str(v))
    return float(m.group(0).replace(",", "")) if m else None


def _rows_to_plans(rows: List[List[Any]]) -> List[dict]:
    rows = [[("" if c is None else str(c).strip()) for c in r] for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = [h.lower() for h in rows[0]]
    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if n in h:
                    return i
        return None
    ci_name, ci_price = col("plan", "name", "tier", "package"), col("price", "amount", "cost", "monthly")
    ci_int, ci_feat, ci_desc = col("interval", "billing", "period"), col("feature", "include", "benefit"), col("desc", "summary")
    has_header = ci_name is not None and ci_price is not None
    body = rows[1:] if has_header else rows
    if not has_header:
        ci_name, ci_price, ci_feat = 0, 1, 2
    plans = []
    for r in body:
        def g(i):
            return r[i] if i is not None and i < len(r) else ""
        name, price = g(ci_name), _to_price(g(ci_price))
        if not name or price is None:
            continue
        interval_raw = (g(ci_int) or "month").lower()
        interval = "year" if "year" in interval_raw or "annual" in interval_raw else ("one_time" if "one" in interval_raw else "month")
        feats = [f.strip(" •-–*") for f in re.split(r"[\n;,|•]", g(ci_feat)) if f.strip(" •-–*")]
        plans.append({"name": name, "price": price, "interval": interval, "features": feats, "description": g(ci_desc)})
    return plans


def parse_pricing_file(filename: str, data: bytes) -> List[dict]:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        plans = []
        for ws in wb.worksheets:
            plans += _rows_to_plans([list(r) for r in ws.iter_rows(values_only=True)])
        return plans
    if ext == "csv":
        return _rows_to_plans(list(csv.reader(io.StringIO(data.decode("utf-8", "ignore")))))
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        plans = []
        for t in d.tables:
            plans += _rows_to_plans([[c.text for c in row.cells] for row in t.rows])
        if not plans:
            # fallback: lines like "Pro - $99/mo - feature; feature"
            for p in d.paragraphs:
                parts = [x.strip() for x in re.split(r"\s[-–:|]\s", p.text) if x.strip()]
                if len(parts) >= 2 and _to_price(parts[1]) is not None:
                    plans += _rows_to_plans([parts])
        return plans
    raise HTTPException(400, "Unsupported file. Upload .xlsx, .csv or .docx")


def sync_plans_to_stripe(owner_id: str, plans: List[dict]) -> List[dict]:
    out = []
    for p in plans:
        pid = f"{owner_id}_{slugify(p['name'])}"
        lookup = f"{pid}_{p['interval']}"
        product = None
        for pr in stripe.Product.list(active=True, limit=100).auto_paging_iter():
            if pr.to_dict().get("metadata", {}).get("emergent_product_id") == pid:
                product = pr
                break
        if not product:
            product = stripe.Product.create(name=p["name"], tax_code="txcd_10103001",
                                            metadata={"managed_by": "emergent", "emergent_product_id": pid})
        amount = int(round(p["price"] * 100))
        existing = stripe.Price.list(lookup_keys=[lookup], active=True, limit=1).data
        if existing and existing[0].unit_amount != amount:
            stripe.Price.modify(existing[0].id, active=False)
            existing = []
        if not existing:
            kw = dict(product=product.id, unit_amount=amount, currency="usd", lookup_key=lookup, transfer_lookup_key=True)
            if p["interval"] in ("month", "year"):
                kw["recurring"] = {"interval": p["interval"]}
            stripe.Price.create(**kw)
        out.append({**p, "plan_id": pid, "lookup_key": lookup, "currency": "usd"})
    return out


# ---------- Registration ----------
def register(api, db, get_current_user, get_user_app, log_activity):

    # ===== AI MEDIA STUDIO =====
    async def _eleven_key() -> str:
        s = await db.settings.find_one({"key": "elevenlabs_api_key"}, {"_id": 0})
        return (s or {}).get("value") or os.environ.get("ELEVENLABS_API_KEY", "")

    @api.get("/media/config")
    async def media_config(user: dict = Depends(get_current_user)):
        key = await _eleven_key()
        return {"elevenlabs": bool(key), "images": bool(EMERGENT_LLM_KEY), "video": bool(EMERGENT_LLM_KEY),
                "video_model": FAL_VIDEO_ENDPOINT, "image_model": "gpt-image-1"}

    @api.post("/media/config/elevenlabs")
    async def set_eleven_key(body: ElevenKeyIn, user: dict = Depends(get_current_user)):
        key = body.api_key.strip()
        if len(key) < 10:
            raise HTTPException(400, "Invalid key")
        try:
            from elevenlabs.client import ElevenLabs
            await asyncio.to_thread(lambda: ElevenLabs(api_key=key).voices.get_all())
        except Exception as e:
            raise HTTPException(400, f"ElevenLabs rejected this key: {str(e)[:120]}")
        await db.settings.update_one({"key": "elevenlabs_api_key"}, {"$set": {"value": key, "updated_by": user["user_id"], "updated_at": now_iso()}}, upsert=True)
        return {"ok": True, "elevenlabs": True}

    @api.get("/media/voices")
    async def list_voices(user: dict = Depends(get_current_user)):
        key = await _eleven_key()
        if not key:
            raise HTTPException(503, "ElevenLabs not configured")
        from elevenlabs.client import ElevenLabs
        res = await asyncio.to_thread(lambda: ElevenLabs(api_key=key).voices.get_all())
        return [{"voice_id": v.voice_id, "name": v.name, "category": getattr(v, "category", None)} for v in res.voices]

    @api.get("/apps/{app_id}/media")
    async def list_media(app_id: str, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        return await db.media_assets.find({"app_id": app_id}, {"_id": 0}).sort("created_at", -1).limit(60).to_list(60)

    @api.delete("/apps/{app_id}/media/{asset_id}")
    async def delete_media(app_id: str, asset_id: str, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        await db.media_assets.delete_one({"app_id": app_id, "asset_id": asset_id})
        return {"ok": True}

    async def _save_asset(app_id, user_id, kind, prompt, **extra):
        doc = {"asset_id": uid("ast"), "app_id": app_id, "user_id": user_id, "kind": kind, "prompt": prompt,
               "created_at": now_iso(), **extra}
        await db.media_assets.insert_one(doc)
        doc.pop("_id", None)
        return doc

    @api.post("/apps/{app_id}/media/image")
    async def gen_image(app_id: str, body: ImageIn, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        if not EMERGENT_LLM_KEY:
            raise HTTPException(500, "LLM key missing")
        from emergentintegrations.llm.openai.image_generation import OpenAIImageGeneration
        gen = OpenAIImageGeneration(api_key=EMERGENT_LLM_KEY)
        prompt = f"{body.prompt}. High-converting {body.style} visual, premium, clean composition, no text overlays."
        try:
            images = await gen.generate_images(prompt=prompt, model="gpt-image-1", number_of_images=1)
        except Exception as e:
            logger.exception("image gen failed")
            raise HTTPException(500, f"Image generation failed: {str(e)[:160]}")
        if not images:
            raise HTTPException(500, "No image generated")
        b64 = base64.b64encode(images[0]).decode()
        doc = await _save_asset(app_id, user["user_id"], "image", body.prompt, data_url=f"data:image/png;base64,{b64}", model="gpt-image-1")
        await log_activity(app_id, user["user_id"], "media.image", f"Generated image: {body.prompt[:60]}")
        return doc

    @api.post("/apps/{app_id}/media/voice")
    async def gen_voice(app_id: str, body: VoiceIn, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        key = await _eleven_key()
        if not key:
            raise HTTPException(503, "ElevenLabs API key not configured. Add it in AI Media → Connect ElevenLabs.")
        from elevenlabs.client import ElevenLabs
        def _run():
            client = ElevenLabs(api_key=key)
            return b"".join(client.text_to_speech.convert(text=body.text, voice_id=body.voice_id, model_id="eleven_multilingual_v2"))
        try:
            audio = await asyncio.to_thread(_run)
        except Exception as e:
            raise HTTPException(500, f"Voice generation failed: {str(e)[:160]}")
        b64 = base64.b64encode(audio).decode()
        doc = await _save_asset(app_id, user["user_id"], "voice", body.text[:200], data_url=f"data:audio/mpeg;base64,{b64}",
                                voice_id=body.voice_id, model="eleven_multilingual_v2")
        await log_activity(app_id, user["user_id"], "media.voice", f"Generated voiceover ({len(body.text)} chars)")
        return doc

    @api.post("/apps/{app_id}/media/video")
    async def gen_video(app_id: str, body: VideoIn, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        if not EMERGENT_LLM_KEY:
            raise HTTPException(500, "LLM key missing")
        task = {"task_id": uid("task"), "app_id": app_id, "user_id": user["user_id"], "kind": "video",
                "status": "running", "prompt": body.prompt, "created_at": now_iso()}
        await db.media_tasks.insert_one(dict(task))
        payload = {"prompt": body.prompt[:2000], "duration": body.duration if body.duration in ("6", "10") else "6", "prompt_optimizer": True}

        async def _worker():
            try:
                result = await asyncio.to_thread(run_fal_video, payload)
                url = (result.get("video") or {}).get("url")
                if not url:
                    raise RuntimeError("No video URL returned")
                doc = await _save_asset(app_id, user["user_id"], "video", body.prompt, url=url, model=FAL_VIDEO_ENDPOINT)
                await db.media_tasks.update_one({"task_id": task["task_id"]}, {"$set": {"status": "done", "asset_id": doc["asset_id"], "url": url}})
                await log_activity(app_id, user["user_id"], "media.video", f"Generated video: {body.prompt[:60]}")
            except Exception as e:
                logger.exception("video gen failed")
                await db.media_tasks.update_one({"task_id": task["task_id"]}, {"$set": {"status": "failed", "error": str(e)[:200]}})

        asyncio.create_task(_worker())
        task.pop("_id", None)
        return task

    @api.get("/apps/{app_id}/media/tasks/{task_id}")
    async def media_task(app_id: str, task_id: str, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        t = await db.media_tasks.find_one({"task_id": task_id, "app_id": app_id}, {"_id": 0})
        if not t:
            raise HTTPException(404, "Task not found")
        return t

    # ===== BILLING =====
    async def _plans_for(owner_id: str) -> List[dict]:
        doc = await db.plan_catalogs.find_one({"owner_id": owner_id}, {"_id": 0})
        if doc:
            return doc["plans"]
        plans = await asyncio.to_thread(sync_plans_to_stripe, owner_id, DEFAULT_PLANS)
        await db.plan_catalogs.update_one({"owner_id": owner_id}, {"$set": {"plans": plans, "source": "default", "updated_at": now_iso()}}, upsert=True)
        return plans

    @api.get("/billing/plans")
    async def get_plans(user: dict = Depends(get_current_user)):
        doc = await db.plan_catalogs.find_one({"owner_id": user["user_id"]}, {"_id": 0})
        plans = await _plans_for(user["user_id"])
        return {"plans": plans, "source": (doc or {}).get("source", "default"), "source_file": (doc or {}).get("source_file"),
                "tax_mode": TAX_MODE}

    @api.post("/billing/plans/import")
    async def import_plans(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
        data = await file.read()
        plans = parse_pricing_file(file.filename or "pricing.xlsx", data)
        if not plans:
            raise HTTPException(400, "No pricing rows detected. Use columns like: Plan, Price, Interval, Features, Description")
        plans = plans[:24]
        synced = await asyncio.to_thread(sync_plans_to_stripe, user["user_id"], plans)
        await db.plan_catalogs.update_one({"owner_id": user["user_id"]},
                                          {"$set": {"plans": synced, "source": "import", "source_file": file.filename, "updated_at": now_iso()}}, upsert=True)
        return {"plans": synced, "count": len(synced), "source_file": file.filename}

    @api.put("/billing/plans")
    async def replace_plans(body: PlansReplaceIn, user: dict = Depends(get_current_user)):
        plans = [p.model_dump() for p in body.plans][:24]
        synced = await asyncio.to_thread(sync_plans_to_stripe, user["user_id"], plans)
        await db.plan_catalogs.update_one({"owner_id": user["user_id"]},
                                          {"$set": {"plans": synced, "source": "manual", "updated_at": now_iso()}}, upsert=True)
        return {"plans": synced}

    @api.post("/billing/checkout")
    async def checkout(body: CheckoutIn, user: dict = Depends(get_current_user)):
        app_doc = await get_user_app(body.app_id, user)
        plans = await _plans_for(app_doc["owner_id"])
        plan = next((p for p in plans if p["lookup_key"] == body.lookup_key), None)
        if not plan:
            raise HTTPException(404, "Plan not found")
        prices = stripe.Price.list(lookup_keys=[body.lookup_key], active=True, limit=1).data
        if not prices:
            raise HTTPException(500, "Stripe price missing")
        price = prices[0]
        kwargs = dict(
            line_items=[{"price": price.id, "quantity": 1}],
            mode="subscription" if price.recurring else "payment",
            success_url=f"{body.origin_url}/payment/success?session_id={{CHECKOUT_SESSION_ID}}",
            cancel_url=f"{body.origin_url}/payment/cancel?app_id={body.app_id}",
            metadata={"user_id": user["user_id"], "app_id": body.app_id, "lookup_key": body.lookup_key, "plan_name": plan["name"]},
        )
        def _create():
            try:
                return stripe.checkout.Session.create(**kwargs, managed_payments={"enabled": True})
            except stripe.error.InvalidRequestError as e:
                msg = (e.user_message or str(e)).lower()
                if "managed payments" in msg or "ineligible" in msg:
                    return stripe.checkout.Session.create(**kwargs, automatic_tax={"enabled": True}, billing_address_collection="required")
                raise
        try:
            session = await asyncio.to_thread(_create)
        except stripe.error.StripeError as e:
            raise HTTPException(500, f"Stripe error: {str(e)[:160]}")
        await db.payment_transactions.insert_one({
            "session_id": session.id, "user_id": user["user_id"], "app_id": body.app_id, "lookup_key": body.lookup_key,
            "plan_name": plan["name"], "amount": float(price.unit_amount or 0) / 100, "currency": price.currency,
            "status": "initiated", "payment_status": "pending", "created_at": now_iso(), "updated_at": now_iso(),
        })
        return {"checkout_url": session.url, "session_id": session.id}

    async def _mark_paid(session_id: str, extra: dict):
        r = await db.payment_transactions.find_one_and_update(
            {"session_id": session_id, "payment_status": {"$ne": "paid"}},
            {"$set": {"status": "completed", "payment_status": "paid", "updated_at": now_iso(), **extra}},
        )
        if r:
            await db.apps.update_one({"app_id": r["app_id"]}, {"$set": {"plan": r["plan_name"], "plan_lookup_key": r["lookup_key"], "billing_status": "active"}})
            await log_activity(r["app_id"], r["user_id"], "billing.paid", f"Subscribed to {r['plan_name']} plan")
            wfs = await db.workflows.find({"app_id": r["app_id"], "trigger": "payment_succeeded", "enabled": True}, {"_id": 0}).to_list(20)
            if wfs:
                from server import WF_HOOKS
                await WF_HOOKS["fire_event"](r["app_id"], "payment_succeeded", {"email": "", "plan": r["plan_name"], "amount": r.get("amount"), "tier": r["plan_name"].lower()})

    @api.get("/billing/status/{session_id}")
    async def billing_status(session_id: str):
        rec = await db.payment_transactions.find_one({"session_id": session_id}, {"_id": 0})
        if not rec:
            raise HTTPException(404, "Transaction not found")
        if rec.get("payment_status") != "paid":
            try:
                s = await asyncio.to_thread(stripe.checkout.Session.retrieve, session_id)
                if s.payment_status == "paid" or s.status == "complete":
                    await _mark_paid(session_id, {"stripe_subscription_id": s.subscription, "stripe_payment_intent_id": s.payment_intent})
                    rec = await db.payment_transactions.find_one({"session_id": session_id}, {"_id": 0})
            except stripe.error.StripeError:
                pass
        return {"session_id": rec["session_id"], "status": rec["status"], "payment_status": rec["payment_status"],
                "app_id": rec.get("app_id"), "plan_name": rec.get("plan_name")}

    @api.post("/stripe/webhook")
    async def stripe_webhook(request: Request):
        payload = await request.body()
        sig = request.headers.get("stripe-signature", "")
        try:
            event = stripe.Webhook.construct_event(payload, sig, STRIPE_WEBHOOK_SECRET)
        except Exception:
            raise HTTPException(400, "Invalid signature")
        obj, t = event["data"]["object"], event["type"]
        if t in ("checkout.session.completed", "checkout.session.async_payment_succeeded"):
            await _mark_paid(obj["id"], {"stripe_subscription_id": obj.get("subscription"), "stripe_payment_intent_id": obj.get("payment_intent")})
        elif t == "checkout.session.async_payment_failed":
            await db.payment_transactions.update_one({"session_id": obj["id"]}, {"$set": {"status": "failed", "payment_status": "failed", "updated_at": now_iso()}})
        elif t == "checkout.session.expired":
            await db.payment_transactions.update_one({"session_id": obj["id"]}, {"$set": {"status": "expired", "payment_status": "expired", "updated_at": now_iso()}})
        return {"status": "ok"}

    @api.get("/billing/transactions")
    async def transactions(app_id: Optional[str] = None, user: dict = Depends(get_current_user)):
        q = {"user_id": user["user_id"]}
        if app_id:
            q["app_id"] = app_id
        return await db.payment_transactions.find(q, {"_id": 0}).sort("created_at", -1).limit(50).to_list(50)

    # ===== CUSTOM DOMAINS =====
    def _dns_check(domain: str, token: str) -> dict:
        res = dns.resolver.Resolver()
        res.lifetime = 5
        out = {"cname": {"expected": DNS_CNAME_TARGET, "found": [], "ok": False},
               "txt": {"name": f"_lucio-verify.{domain}", "expected": token, "found": [], "ok": False},
               "a": {"found": []}}
        try:
            out["cname"]["found"] = [str(r.target).rstrip(".") for r in res.resolve(domain, "CNAME")]
        except Exception:
            pass
        try:
            out["a"]["found"] = [str(r) for r in res.resolve(domain, "A")]
        except Exception:
            pass
        try:
            out["txt"]["found"] = [b"".join(r.strings).decode("utf-8", "ignore") for r in res.resolve(f"_lucio-verify.{domain}", "TXT")]
        except Exception:
            pass
        out["cname"]["ok"] = DNS_CNAME_TARGET in out["cname"]["found"]
        out["txt"]["ok"] = token in out["txt"]["found"]
        out["resolves"] = bool(out["cname"]["found"] or out["a"]["found"])
        return out

    @api.post("/apps/{app_id}/domain")
    async def set_domain(app_id: str, body: DomainIn, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        domain = body.domain.strip().lower().removeprefix("https://").removeprefix("http://").rstrip("/")
        if not re.match(r"^(?=.{4,253}$)([a-z0-9-]+\.)+[a-z]{2,}$", domain):
            raise HTTPException(400, "Enter a valid domain like app.client.com")
        token = f"lucio-verify={uid('tok')}"
        await db.apps.update_one({"app_id": app_id}, {"$set": {"custom_domain": domain, "domain_token": token, "domain_status": "pending",
                                                            "domain_checked_at": None, "domain_dns": None}})
        await log_activity(app_id, user["user_id"], "domain.set", f"Custom domain set: {domain}")
        return await db.apps.find_one({"app_id": app_id}, {"_id": 0})

    @api.delete("/apps/{app_id}/domain")
    async def remove_domain(app_id: str, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        await db.apps.update_one({"app_id": app_id}, {"$unset": {"custom_domain": "", "domain_token": "", "domain_status": "", "domain_dns": "", "domain_checked_at": ""}})
        return await db.apps.find_one({"app_id": app_id}, {"_id": 0})

    @api.post("/apps/{app_id}/domain/verify")
    async def verify_domain(app_id: str, user: dict = Depends(get_current_user)):
        doc = await get_user_app(app_id, user)
        if not doc.get("custom_domain"):
            raise HTTPException(400, "No domain set")
        dns_res = await asyncio.to_thread(_dns_check, doc["custom_domain"], doc.get("domain_token", ""))
        status_v = "verified" if dns_res["cname"]["ok"] and dns_res["txt"]["ok"] else ("partial" if dns_res["cname"]["ok"] or dns_res["txt"]["ok"] else "pending")
        await db.apps.update_one({"app_id": app_id}, {"$set": {"domain_status": status_v, "domain_dns": dns_res, "domain_checked_at": now_iso()}})
        await log_activity(app_id, user["user_id"], "domain.verify", f"DNS check for {doc['custom_domain']}: {status_v}",
                           "info" if status_v == "verified" else "warning")
        return await db.apps.find_one({"app_id": app_id}, {"_id": 0})

    # ===== LIVE PREVIEW LINK =====
    @api.post("/apps/{app_id}/preview/regenerate")
    async def regen_preview(app_id: str, user: dict = Depends(get_current_user)):
        await get_user_app(app_id, user)
        token = uid("pv") + uid("")[1:]
        await db.apps.update_one({"app_id": app_id}, {"$set": {"preview_token": token, "preview_enabled": True}})
        await log_activity(app_id, user["user_id"], "preview.link", "Live preview link regenerated")
        return await db.apps.find_one({"app_id": app_id}, {"_id": 0})

    @api.post("/apps/{app_id}/preview/toggle")
    async def toggle_preview(app_id: str, user: dict = Depends(get_current_user)):
        doc = await get_user_app(app_id, user)
        enabled = not doc.get("preview_enabled", False)
        upd = {"preview_enabled": enabled}
        if enabled and not doc.get("preview_token"):
            upd["preview_token"] = uid("pv") + uid("")[1:]
        await db.apps.update_one({"app_id": app_id}, {"$set": upd})
        return await db.apps.find_one({"app_id": app_id}, {"_id": 0})

    @api.get("/public/preview/{token}")
    async def public_preview(token: str):
        doc = await db.apps.find_one({"preview_token": token, "preview_enabled": True}, {"_id": 0})
        if not doc:
            raise HTTPException(404, "Preview link is invalid or has been revoked")
        page = await db.pages.find_one({"app_id": doc["app_id"]}, {"_id": 0})
        return {"app": {k: doc.get(k) for k in ("name", "industry", "description", "color", "status", "custom_domain", "video_url")},
                "blocks": (page or {}).get("blocks", []), "updated_at": (page or {}).get("updated_at")}
